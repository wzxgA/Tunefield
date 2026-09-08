"""T2 · 结构化解析：txt/md 结构、pdf/docx 抽取、code 函数边界、逐文件扫描。"""

from __future__ import annotations

import io

from tunefield import config
from tunefield.pipeline.parsers import parse_bytes, scan_dataset_files


def test_txt_paragraphs():
    data = "第一段内容。\n\n第二段内容。\n".encode("utf-8")
    r = parse_bytes("a.txt", data)
    assert r["kind"] == "txt" and r["status"] == "ok"
    assert r["blocks"] == 2
    assert "第一段内容" in r["segments"][0]["text"]
    assert r["segments"][1]["meta"]["kind"] == "paragraph"


def test_txt_gbk_decode():
    data = "中文乱码测试。\n".encode("gb18030")
    r = parse_bytes("c.txt", data)
    assert r["status"] == "ok"
    assert "中文乱码测试" in r["segments"][0]["text"]


def test_txt_empty_and_unsupported():
    assert parse_bytes("e.txt", b"")["status"] == "empty"
    r = parse_bytes("x.bin", b"\x00\x01\x02")
    assert r["status"] == "unsupported"
    assert r["kind"] == "other"


def test_md_headings_and_fence():
    src = (
        "# 标题一\n"
        "正文甲\n\n"
        "```python\n"
        "x = 1\n"
        "```\n\n"
        "## 子标题\n"
        "正文乙\n"
    )
    r = parse_bytes("d.md", src.encode("utf-8"))
    assert r["status"] == "ok" and r["kind"] == "md"
    metas = [s["meta"] for s in r["segments"]]

    headings = [m for m in metas if m["kind"] == "heading"]
    assert [h["title"] for h in headings] == ["标题一", "子标题"]
    assert headings[1]["level"] == 2
    assert headings[1]["ancestors"] == ["标题一"]  # 子标题继承祖先链

    code_meta = next(m for m in metas if m["kind"] == "code")
    assert code_meta["fence_lang"] == "python"

    # 「子标题」下的正文段完整继承标题链：["标题一", "子标题"]
    deep_para = next(
        m for m in metas
        if m["kind"] == "paragraph" and m["ancestors"] == ["标题一", "子标题"]
    )
    assert deep_para is not None


def test_pdf_pages():
    import pymupdf  # PyMuPDF（中文渲染需外部字体，样例用英文验证文本层抽取）

    doc = pymupdf.open()
    for text in ("First page text", "Second page text"):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    buf = doc.tobytes()
    doc.close()

    r = parse_bytes("doc.pdf", buf)
    assert r["kind"] == "pdf" and r["status"] == "ok"
    pages = [s["meta"]["page"] for s in r["segments"]]
    assert pages == [1, 2]
    assert "Second page text" in r["segments"][1]["text"]


def test_pdf_corrupt():
    r = parse_bytes("bad.pdf", b"not a real pdf at all")
    assert r["status"] == "error"
    assert "打开失败" in r["error"]


def test_docx_paragraph_and_table():
    import docx

    document = docx.Document()
    document.add_paragraph("第一段 hello")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    buf = io.BytesIO()
    document.save(buf)

    r = parse_bytes("d.docx", buf.getvalue())
    assert r["kind"] == "docx" and r["status"] == "ok"
    metas = [s["meta"] for s in r["segments"]]
    assert any(m["kind"] == "paragraph" for m in metas)
    table_seg = next(s for s in r["segments"] if s["meta"]["kind"] == "table")
    assert table_seg["text"] == "A | B"


def test_code_python_defs():
    src = (
        "import os\n\n"
        "class Greeter:\n"
        "    def greet(self, name):\n"
        '        return f"hi {name}"\n'
        "\n"
        "def main():\n"
        '    g = Greeter()\n'
        '    print(g.greet("x"))\n'
    )
    r = parse_bytes("mod.py", src.encode("utf-8"))
    assert r["kind"] == "code"
    assert r["language"] == "python"
    assert r["status"] == "ok"
    assert r["structure"] == "tree-sitter"

    # 取最外层定义整体：Greeter（含内部方法）与 main，不把方法单独切出
    symbols = [s["meta"]["symbol"] for s in r["segments"]]
    assert symbols == ["Greeter", "main"]
    greeter = r["segments"][0]
    assert greeter["meta"]["definition"] == "class_definition"
    assert greeter["text"].startswith("class Greeter")
    assert "def greet" in greeter["text"]  # 类整体一块，不腰斩


def test_code_typescript_and_rust_smoke():
    ts = parse_bytes("a.ts", b"interface A {\n  x: number\n}\n")
    assert ts["language"] == "typescript" and ts["status"] == "ok"
    rs = parse_bytes("a.rs", b"fn main() {\n    println!(\"hi\");\n}\n")
    assert rs["language"] == "rust" and rs["status"] == "ok"
    assert rs["segments"][0]["meta"]["symbol"] == "main"


def test_scan_dataset_files(tmp_path, monkeypatch):
    raw = tmp_path / "raw" / "hash1"
    (raw / "sub").mkdir(parents=True)
    (raw / "a.py").write_text("def hello():\n    return 1\n", encoding="utf-8")
    (raw / "b.txt").write_text("正文内容\n", encoding="utf-8")
    (raw / "sub" / "c.bin").write_bytes(b"\x00\x01")
    monkeypatch.setattr(config, "RAW_DIR", tmp_path / "raw")

    files = scan_dataset_files({"content_hash": "hash1"})
    assert {f["name"] for f in files} == {"a.py", "b.txt", "sub/c.bin"}
    by_name = {f["name"]: f for f in files}
    assert by_name["a.py"]["status"] == "ok"
    assert by_name["a.py"]["language"] == "python"
    assert "def hello" in by_name["a.py"]["preview"]
    assert by_name["sub/c.bin"]["status"] == "unsupported"
    assert len(by_name["b.txt"]["preview"]) > 0
