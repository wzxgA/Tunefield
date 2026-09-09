"""T12 · 生成三形态领域样例数据（纯文本/文档/代码），供验收与双通道核对。

文件**平铺**在一个目录（与 Web 上传的文件集合语义一致；CLI 目录接入同样平铺），
文件名带形态前缀。样例量刻意小巧（单测与双通道核对用）；真机 8GB 卡验收可在
此目录继续追加真实语料后跑全流程。目录内含一个损坏 PDF，验证解析错误态不阻断。
"""

from __future__ import annotations

from pathlib import Path


def _seed() -> str:
    return "Tunefield 领域专属小模型训练平台验收语料。专注中文领域文本的解析、清洗、切片与训练体验。\n"


def _write_text_files(root: Path) -> None:
    (root / "txt-intro.txt").write_text(
        "".join(_seed() for _ in range(120)), encoding="utf-8"
    )
    (root / "txt-gbk_note.txt").write_text(
        "中文编码验证：这是一段 GBK 编码的说明文本。" * 60, encoding="gbk"
    )
    md = ["# 纯文本领域手册\n\n", "## 使用说明\n\n"]
    md.append("".join(_seed() for _ in range(80)))
    (root / "md-guide.md").write_text("".join(md), encoding="utf-8")


def _write_doc_files(root: Path) -> None:
    import pymupdf

    pdf = pymupdf.open()
    for no in range(1, 5):
        page = pdf.new_page()
        page.insert_text((72, 72), f"Tunefield document page {no}")
        page.insert_text((72, 96), "This is acceptance pdf text for pipeline.")
    pdf.save(str(root / "pdf-spec.pdf"))
    pdf.close()

    import docx

    document = docx.Document()
    document.add_heading("Tunefield DOCX Acceptance", 0)
    for _ in range(40):
        document.add_paragraph("Docx paragraph for pipeline acceptance test.")
    table = document.add_table(rows=3, cols=2)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"cell-{i}-{j}"
    document.save(str(root / "docx-notes.docx"))

    # 坏文件：损坏的 PDF（应标解析失败且不阻断整体）
    (root / "broken.pdf").write_bytes(b"%PDF-1.7 broken-not-a-real-pdf\x00\x01")


def _write_code_files(root: Path) -> None:
    py_lines = []
    for idx in range(1, 6):
        py_lines.append(f"def helper_{idx}(x: int) -> int:\n    return x * {idx}\n\n")
    py_lines.append("class Trainer:\n    def __init__(self):\n        self.loss = 0.0\n\n")
    py_lines.append("    def fit(self, epochs: int) -> None:\n        for _ in range(epochs):\n            self.loss += 0.1\n\n")
    (root / "py-trainer.py").write_text("".join(py_lines), encoding="utf-8")

    js = ["export function buildModel(name) {\n  return { name };\n}\n\n"]
    js.append("class Dataset {\n  load(path) {\n    this.path = path;\n  }\n}\n\n")
    (root / "js-dataset.js").write_text("".join(js), encoding="utf-8")


def build_corpus(root: Path) -> Path:
    """在 root 下生成平铺的三形态样例（文件与目录同级），返回 corpus 目录。"""
    root.mkdir(parents=True, exist_ok=True)
    _write_text_files(root)
    _write_doc_files(root)
    _write_code_files(root)
    return root


if __name__ == "__main__":  # pragma: no cover
    import sys

    target = Path(sys.argv[1] if len(sys.argv) > 1 else "t12_corpus")
    build_corpus(target)
    print(f"已生成验收样例：{target}")
